"""Focused regressions for spreadsheet, nutrition scan, and focus timer."""
import importlib.util
import sys
import tempfile
import threading
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from plugins import excel_writer, pomodoro


class SpreadsheetTests(unittest.TestCase):
    def test_workbook_total_and_literal_text(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "book.xlsx"
            info = excel_writer._build_xlsx({
                "title": "School/expenses", "headers": ["Item", "Amount"],
                "rows": [["=HYPERLINK(\"https://example.com\")", 12], ["Food", 3]],
                "total_columns": [1],
            }, path)
            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual(ws.title, "Schoolexpenses")
            self.assertEqual(ws["A2"].data_type, "s")
            self.assertEqual(ws["B4"].value, "=SUM(B2:B3)")
            self.assertEqual(ws["A4"].value, "TOTAL")
            self.assertTrue(info["total"])

    def test_two_requests_do_not_overwrite(self):
        with tempfile.TemporaryDirectory() as td:
            spec = {"headers": ["One"], "rows": [[1]]}
            with patch.object(excel_writer, "_output_dir", return_value=Path(td)), \
                 patch.object(excel_writer, "_build_spec", return_value=spec):
                excel_writer.run({"request": "one"})
                excel_writer.run({"request": "one"})
            self.assertEqual(len(list(Path(td).glob("*.xlsx"))), 2)

    def test_bad_model_spec_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaises(ValueError):
                excel_writer._build_xlsx({"headers": ["A"], "rows": "bad"}, Path(td) / "bad.xlsx")


class TimerTests(unittest.TestCase):
    def test_stop_then_start_waits_for_previous_worker(self):
        gate = threading.Event()
        entered = threading.Event()
        original = pomodoro._worker

        def worker(_player, _stop):
            entered.set()
            gate.wait(2)
            original(_player, _stop)

        with patch.object(pomodoro, "_worker", worker):
            try:
                self.assertIn("Starting", pomodoro.run({"action": "start"}))
                self.assertTrue(entered.wait(1))
                self.assertIn("stopped", pomodoro.run({"action": "stop"}))
                self.assertIn("already running", pomodoro.run({"action": "start"}))
            finally:
                gate.set()
                pomodoro._state["thread"].join(2)
        self.assertFalse(pomodoro._state["running"])

    def test_phase_stops_without_recording(self):
        stop = threading.Event()
        stop.set()
        self.assertFalse(pomodoro._run_phase(None, stop, "work", 1, ""))
        self.assertIn("start, stop, status, or stats", pomodoro.run({"action": "typo"}))


class NutritionTests(unittest.TestCase):
    def test_camera_released_and_lock_reset_on_analysis_error(self):
        # The runner does not have OpenCV; these stubs let us check the device lifecycle.
        numpy_stub = types.ModuleType("numpy")
        numpy_stub.ndarray = object
        with patch.dict(sys.modules, {"cv2": types.ModuleType("cv2"),
                                      "numpy": numpy_stub}):
            spec = importlib.util.spec_from_file_location(
                "calorie_test", Path(__file__).resolve().parents[1] / "plugins" / "calorie_counter (4).py")
            mod = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(mod)
        class Camera:
            released = False
            def read(self):
                return True, object()
            def release(self):
                self.released = True
        cam = Camera()
        with patch("memory.config_manager.get_gemini_key", return_value="test-key"), \
             patch.object(mod, "_open_camera", return_value=cam), \
             patch.object(mod, "_analyze", side_effect=RuntimeError("offline")), \
             patch.object(mod, "_LIVE_SCAN_SECONDS", 0):
            result = mod.run({"query": "Calories?"})
        self.assertIn("offline", result)
        self.assertTrue(cam.released)
        self.assertTrue(mod._camera_lock.acquire(blocking=False))
        mod._camera_lock.release()


if __name__ == "__main__":
    unittest.main()
